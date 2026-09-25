#!/usr/bin/env python3
"""Build Supplementary_Tables.xlsx (Supp. Tables S1 to S12) from the workflow outputs.

    python workflow/scripts/make_supplement.py [out.xlsx]

Every value is read from a workflow output; nothing is typed by hand except
labels and captions, whose counts and software versions are also read from
the outputs (results/manuscript_values.tsv for versions). The summed length
of the three chromosome scale sequences falls back on V4 measurements of the
synteny FASTA files when those files are absent. Each sheet names the files
it was built from. S12 (the outgroup analyses, V5) is written when
results/phylo/rooted and results/phylo/dstat hold their outputs. The
outgroup is read from config/samples.tsv (or CX_SAMPLE_SHEET).
"""
import collections
import csv
import glob
import json
import math
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import forms as samples_meta  # noqa: E402

REPO = samples_meta.REPO
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, "tables", "Supplementary_Tables.xlsx")
RELEASE = "v3.0"          # the git tag of the V5 workflow cited in the README sheet

ORDER = samples_meta.INGROUP
OUTGROUP = samples_meta.outgroup()
ALL5 = ORDER + [OUTGROUP]
FORM = samples_meta.FORM
ACCESSION = samples_meta.ACCESSION
OG = FORM[OUTGROUP]                      # the outgroup's name in captions
CHROM_BP_V4 = {"Cx_quinquefasciatus": 559588684, "Cx_pallens": 549446285,
               "Cx_molestus": 530551183, "Cx_pipiens": 532519368}
CHROM_BP = {}
for _s in ALL5:
    _bp = samples_meta.chromosome_scale_bp(_s)
    if _bp is None:
        _bp = CHROM_BP_V4.get(_s)
    if _bp:
        CHROM_BP[_s] = _bp

FONT = "Arial"
F_TITLE = Font(name=FONT, size=12, bold=True)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_BODY = Font(name=FONT, size=10)
F_NOTE = Font(name=FONT, size=9, italic=True)
THIN = Side(style="thin", color="000000")
NUMFMT = {"int": "#,##0", "pct1": "0.0", "pct2": "0.00", "pct3": "0.000", "f3": "0.000",
          "f4": "0.0000", "f5": "0.00000", "p": "0.0E+00"}


# ------------------------------------------------------------------ helpers
def path(p):
    return os.path.join(REPO, p)


def tsv(p, skip_comments=True):
    with open(path(p)) as fh:
        lines = [l for l in fh if l.strip() and not (skip_comments and l.startswith("#"))]
    return list(csv.DictReader(lines, delimiter="\t"))


def _values():
    p = os.path.join(REPO, "results", "manuscript_values.tsv")
    if not os.path.exists(p):
        return {}
    with open(p) as fh:
        return {(r["section"], r["item"], r["sample"]): r["value"]
                for r in csv.DictReader(fh, delimiter="\t")}


VALUES = _values()


def version(tool, fallback):
    """Version a run used, from results/manuscript_values.tsv (conda
    environments first, then what the outputs record), else the fallback."""
    for sec in ("tools", "tools_observed"):
        v = VALUES.get((sec, tool, ""))
        if v and v != "NA" and not v.startswith("not present"):
            return v.split(",")[0]
    return fallback


def pct_range(values, digits=0):
    lo, hi = min(values), max(values)
    a, b = f"{lo:.{digits}f}", f"{hi:.{digits}f}"
    return f"{a}%" if a == b else f"{a} to {b}%"


def forms(text):
    """Replace workflow sample names with the form names used in the paper."""
    if text is None:
        return None
    s = str(text)
    for k in sorted(FORM, key=len, reverse=True):
        s = s.replace(k, FORM[k])
    return s


def num(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return x
    s = str(x).strip().rstrip("%")
    if s == "" or s.upper() in ("NA", "NAN"):
        return None
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return str(x)


TX = re.compile(r"^(rna-XM_\d{9})(\d)$")


def tx(ids):
    """Restore the version period that the protein FASTA headers drop
    (rna-XM_0382506301 -> rna-XM_038250630.1)."""
    if not ids:
        return ids
    return ", ".join(TX.sub(r"\1.\2", g.strip()) for g in ids.split(","))


def pair_name(p):
    return forms(p).replace("_vs_", " vs ")


class Block:
    def __init__(self, title, columns, rows, notes=None, autofilter=False):
        self.title = title            # e.g. "(a) Gene trees per topology"
        self.columns = columns        # [(key, header, fmt)], fmt in NUMFMT or "text"
        self.rows = rows              # list of dicts
        self.notes = notes or []
        self.autofilter = autofilter


def cell_value(v, fmt):
    if fmt == "text":
        return forms(v) if isinstance(v, str) else v
    if fmt == "bool":
        return {"True": "yes", "False": "no", True: "yes", False: "no"}.get(v, v)
    v = num(v)
    if isinstance(v, float) and math.isnan(v):
        return None
    return v


def display_len(v, fmt):
    if v is None:
        return 0
    if isinstance(v, str):
        return len(v)
    if fmt == "int":
        return len(f"{v:,.0f}")
    if fmt == "p":
        return 8
    if fmt in NUMFMT:
        dec = NUMFMT[fmt].count("0") - 1
        return len(f"{v:.{max(dec, 0)}f}")
    return len(str(v))


def wrap_lines(text, width):
    """Lines taken by text in a column of the given width (greedy word wrap;
    bold Arial 10 fits about 0.85 characters per width unit)."""
    per_line = max(int(width * 0.85), 4)
    lines, cur = 1, 0
    for w in str(text).split():
        need = len(w) if cur == 0 else cur + 1 + len(w)
        if need <= per_line:
            cur = need
        else:
            lines += 1
            cur = len(w)
    return lines


def write_sheet(wb, name, title, caption, source, blocks, freeze=False, max_width=48):
    ws = wb.create_sheet(name)
    ncols = max(len(b.columns) for b in blocks)
    # column widths from headers and data
    widths = [10] * ncols
    for b in blocks:
        for i, (key, head, fmt) in enumerate(b.columns):
            longest_word = max(len(w) for w in head.split())
            widths[i] = max(widths[i], longest_word + 2, min(len(head), 14) + 2)
            for r in b.rows[:3000]:
                widths[i] = max(widths[i], display_len(cell_value(r.get(key), fmt), fmt) + 2)
    widths = [min(w, max_width) for w in widths]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # caption spans enough columns for about 120 characters per line
    span, total = 0, 0
    while span < len(widths) and total < 120:
        total += widths[span]
        span += 1
    span = max(span, 1)
    chars_per_line = max(total * 1.05, 40)

    def text_row(r, text, font, height_lines=None):
        ws.cell(row=r, column=1, value=text).font = font
        if span > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        lines = height_lines or max(1, math.ceil(len(text) / chars_per_line))
        ws.row_dimensions[r].height = 13.0 * lines + 3
        return r + 1

    r = text_row(1, title, F_TITLE, 1)
    ws.row_dimensions[1].height = 18
    r = text_row(r, caption, F_BODY)
    r = text_row(r, "Source: " + source, F_NOTE)
    r += 1
    header_rows = []
    for b in blocks:
        if b.title:
            ws.cell(row=r, column=1, value=b.title).font = F_BOLD
            r += 1
        # header
        hr = r
        header_rows.append(hr)
        lines = 1
        for i, (key, head, fmt) in enumerate(b.columns, 1):
            c = ws.cell(row=r, column=i, value=head)
            c.font = F_BOLD
            c.alignment = Alignment(wrap_text=True, vertical="bottom",
                                    horizontal="left" if fmt in ("text", "bool") else "right")
            c.border = Border(top=THIN, bottom=THIN)
            lines = max(lines, wrap_lines(head, widths[i - 1]))
        ws.row_dimensions[r].height = 13.0 * lines + 3
        r += 1
        first = r
        for row in b.rows:
            for i, (key, head, fmt) in enumerate(b.columns, 1):
                v = cell_value(row.get(key), fmt)
                c = ws.cell(row=r, column=i, value=v)
                c.font = F_BODY
                if isinstance(v, (int, float)) and fmt in NUMFMT:
                    c.number_format = NUMFMT[fmt]
                if fmt in ("text", "bool"):
                    c.alignment = Alignment(horizontal="left")
            r += 1
        last = r - 1
        for i in range(1, len(b.columns) + 1):
            ws.cell(row=last, column=i).border = Border(bottom=THIN)
        if b.autofilter:
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(b.columns))}{last}"
        for n in b.notes:
            r = text_row(r, n, F_NOTE)
        r += 1
    if freeze and len(blocks) == 1:
        ws.freeze_panes = ws.cell(row=header_rows[0] + 1, column=2)
    return ws, sum(len(b.rows) for b in blocks)


# ------------------------------------------------------------------ S1
def s1():
    quast = {r["Assembly"]: r for r in tsv("results/quast/report.tsv", skip_comments=False)}
    asm, gen, prot = [], [], []
    for s in ALL5:
        bj = json.load(open(glob.glob(path(f"results/busco/{s}/short_summary*.json"))[0]))["results"]
        pj = json.load(open(glob.glob(path(f"results/busco_proteins/{s}/short_summary*.json"))[0]))["results"]
        total = int(quast["Total length (>= 0 bp)"][s])
        asm.append({"form": FORM[s], "acc": ACCESSION[s][0], "name": ACCESSION[s][1],
                    "length": total, "seqs": quast["# contigs (>= 0 bp)"][s],
                    "n50": quast["N50"][s], "l50": quast["L50"][s],
                    "largest": quast["Largest contig"][s], "gc": quast["GC (%)"][s],
                    "contigs": bj["Number of contigs"], "cn50": bj["Contigs N50"],
                    "gaps": bj["Percent gaps"],
                    "chrom": CHROM_BP.get(s),
                    "chrom_pct": round(100 * CHROM_BP[s] / total, 1) if s in CHROM_BP else None})
        for d, j in ((gen, bj), (prot, pj)):
            d.append({"form": FORM[s], "n": j["n_markers"],
                      "c": j["Complete BUSCOs"], "s": j["Single copy BUSCOs"],
                      "d": j["Multi copy BUSCOs"], "f": j["Fragmented BUSCOs"],
                      "m": j["Missing BUSCOs"], "cp": j["Complete percentage"],
                      "sp": j["Single copy percentage"], "dp": j["Multi copy percentage"],
                      "fp": j["Fragmented percentage"], "mp": j["Missing percentage"],
                      "e": j.get("internal_stop_codon_count"),
                      "ep": j.get("internal_stop_codon_percent")})
    busco_cols = [("form", "Form", "text"), ("n", "Genes searched", "int"),
                  ("c", "Complete", "int"), ("s", "Complete, single copy", "int"),
                  ("d", "Complete, duplicated", "int"), ("f", "Fragmented", "int"),
                  ("m", "Missing", "int"), ("cp", "Complete (%)", "pct1"),
                  ("sp", "Single copy (%)", "pct1"), ("dp", "Duplicated (%)", "pct1"),
                  ("fp", "Fragmented (%)", "pct1"), ("mp", "Missing (%)", "pct1")]
    blocks = [
        Block("(a) Assembly statistics",
              [("form", "Form", "text"), ("acc", "Accession", "text"),
               ("name", "Assembly", "text"), ("length", "Total length (bp)", "int"),
               ("seqs", "Sequences", "int"), ("n50", "Scaffold N50 (bp)", "int"),
               ("l50", "L50", "int"), ("largest", "Largest sequence (bp)", "int"),
               ("gc", "GC (%)", "pct2"), ("contigs", "Contigs", "int"),
               ("cn50", "Contig N50 (bp)", "int"), ("gaps", "Gaps, N (%)", "pct3"),
               ("chrom", "Three chromosome scale sequences (bp)", "int"),
               ("chrom_pct", "In three chromosome scale sequences (%)", "pct1")], asm),
        Block("(b) BUSCO, genome mode (miniprot)",
              busco_cols + [("e", "Complete with an internal stop codon", "int"),
                            ("ep", "Complete with an internal stop codon (%)", "pct1")], gen),
        Block("(c) BUSCO, protein mode (gene sets used for orthology, one protein per gene)",
              busco_cols, prot),
    ]
    ref_json = json.load(open(glob.glob(path(f"results/busco/{ORDER[0]}/short_summary*.json"))[0]))
    busco_v = ref_json.get("versions", {}).get("busco", "6.1.0")
    lineage = ref_json.get("lineage_dataset", {}).get("name", "diptera_odb10")
    stops = [float(g["ep"]) for g in gen if g.get("ep") not in (None, "")]
    stop_text = (f"; this share ranged from {pct_range(stops, 1)} across the five assemblies"
                 if stops else "")
    chrom_forms = [FORM[s] for s in ALL5 if s in CHROM_BP]
    return ("S1 Assemblies",
            "Supp. Table S1. Assembly statistics and BUSCO completeness",
            f"Statistics of the five assemblies (a) and BUSCO {busco_v} completeness against the "
            f"{lineage} dataset in genome mode (b) and in protein mode on the gene sets used "
            "for orthology (c). Total length, sequences, scaffold N50, L50, largest sequence and "
            f"GC content are from QUAST {version('quast', '5.3.0')} counting all sequences; "
            "contigs, contig N50 and gaps are as reported by BUSCO, which counts scaffolds split "
            "at gaps. The chromosome scale columns give the summed length of the three longest "
            f"sequences ({', '.join(chrom_forms)}). In genome mode, BUSCO also reports the "
            "complete genes whose miniprot alignment contains an internal stop codon"
            f"{stop_text}.",
            "results/quast/report.tsv; results/busco/<form>/short_summary*.json; "
            "results/busco_proteins/<form>/short_summary*.json; "
            "results/synteny/genomes/<form>.chromosomes.fasta; resources/genomes/<form>.fasta",
            blocks)


# ------------------------------------------------------------------ S2
def s2():
    tq = tsv("results/annotation/transfer_quality.tsv")
    rows = []
    for r in sorted(tq, key=lambda x: ALL5.index(x["sample"])):
        d = dict(r)
        d["sample"] = FORM[r["sample"]]
        d["liftoff_flags"] = "no (reference annotation)" if r["sample"] == "Cx_quinquefasciatus" \
            else r["liftoff_flags"]
        rows.append(d)
    by = {r["sample"]: r for r in tq}
    ref_row = by[ORDER[0]]
    # share of the transfers of a partial or exception reference model that fail the ORF check
    fail = []
    for s in ORDER[1:]:
        r = by.get(s, {})
        try:
            bad = int(r["n_invalid_orf"]) - int(r["n_invalid_orf_clean_ref"])
            fail.append(100.0 * bad / int(r["n_ref_model_not_clean"]))
        except (KeyError, ValueError, ZeroDivisionError):
            pass
    og_flags = by.get(OUTGROUP, {}).get("liftoff_flags", "")
    og_text = (f" The outgroup, {OG}, keeps {samples_meta.OUTGROUP_ANNOTATION.get(OUTGROUP, 'its own gene set')} "
               "and carries no Liftoff flags either." if og_flags.startswith("no") else "")
    comp_order = ["core", "shell", "cloud", "outgroup_only", "unassigned"]
    bc = tsv("results/annotation/transfer_quality_by_compartment.tsv")
    bc.sort(key=lambda x: (ALL5.index(x["sample"]), comp_order.index(x["compartment"])))
    for r in bc:
        r["compartment"] = r["compartment"].replace("_", " ")
    blocks = [
        Block("(a) By form",
              [("sample", "Form", "text"), ("n_kept_models", "Kept models", "int"),
               ("liftoff_flags", "Liftoff flags", "text"),
               ("n_invalid_orf", "Without a valid ORF", "int"),
               ("pct_invalid_orf", "Without a valid ORF (%)", "pct2"),
               ("n_inframe_stop", "Internal stop codon", "int"),
               ("pct_inframe_stop", "Internal stop codon (%)", "pct2"),
               ("n_missing_start", "Missing start codon", "int"),
               ("pct_missing_start", "Missing start codon (%)", "pct2"),
               ("n_missing_stop", "Missing stop codon", "int"),
               ("pct_missing_stop", "Missing stop codon (%)", "pct2"),
               ("n_mismatch_ref_protein", "Protein differs from reference", "int"),
               ("pct_mismatch_ref_protein", "Protein differs from reference (%)", "pct2"),
               ("n_ref_model_not_clean", "Source reference model partial or with exception", "int"),
               ("n_clean_ref_models", "Models with a clean source model", "int"),
               ("n_invalid_orf_clean_ref", "Without a valid ORF, clean source model", "int"),
               ("pct_invalid_orf_clean_ref", "Without a valid ORF, clean source model (%)", "pct2")],
              rows),
        Block("(b) By pangenome compartment",
              [("sample", "Form", "text"), ("compartment", "Compartment", "text"),
               ("n_models", "Kept models", "int"), ("n_invalid_orf", "Without a valid ORF", "int"),
               ("pct_invalid_orf", "Without a valid ORF (%)", "pct2")], bc),
    ]
    return ("S2 Transfer quality",
            "Supp. Table S2. Quality of the transferred gene models",
            f"Liftoff {version('liftoff', '1.6.3')} flags for the gene models kept for analysis (one per gene), by form (a) "
            "and by pangenome compartment (b). A valid open reading frame (ORF) has a start codon, "
            "a stop codon and no internal stop codon. 'Protein differs from reference' counts "
            "models whose protein is not identical to the Cx. quinquefasciatus reference protein, "
            "which includes true amino acid differences between forms. Source reference models "
            f"that are partial or carry a RefSeq sequence exception ({int(ref_row['n_ref_model_not_clean']):,} "
            f"of the {int(ref_row['n_kept_models']):,} reference models) usually fail the ORF check in "
            f"any genome ({pct_range(fail)} of their transfers in the ingroup genomes); the last "
            "columns of (a) set their transfers aside. The reference annotation was not "
            "transferred and carries no Liftoff flags; for it, the column lists how many of its "
            f"own models are partial or carry an exception.{og_text} In (b), compartments are "
            "those of the orthogroup holding each model; unassigned models were placed in no "
            "orthogroup.",
            "results/annotation/transfer_quality.tsv; "
            "results/annotation/transfer_quality_by_compartment.tsv",
            blocks)


# ------------------------------------------------------------------ S3
def s3():
    part = tsv("results/pangenome/partitioned_orthogroups.tsv")
    for r in part:
        r["compartment"] = r["compartment"].replace("_", " ")
    cols = [("Orthogroup", "Orthogroup", "text")] + \
           [(s, FORM[s], "int") for s in ALL5] + \
           [("ingroup_count", "Ingroup forms present", "int"),
            ("compartment", "Compartment", "text")]
    return ("S3 Partition",
            "Supp. Table S3. Pangenome partition of the orthogroups",
            f"Number of genes (one protein per gene) of each form in each of the {len(part):,} "
            f"OrthoFinder {version('orthofinder', '3.1.5')} orthogroups, and the compartment "
            "assigned from the four ingroup forms: core (genes from all four), shell (two or "
            f"three), cloud (one) and outgroup only ({OG} genes only). Genes that OrthoFinder "
            "assigned to no orthogroup are not listed.",
            "results/pangenome/partitioned_orthogroups.tsv",
            [Block(None, cols, part, autofilter=True)])


# ------------------------------------------------------------------ S4
def s4():
    summ = tsv("results/pangenome/cloud_composition_summary.tsv")
    for r in summ:
        r["form"] = "all forms" if r["form"] == "total" else r["form"]
    per = tsv("results/pangenome/cloud_composition.tsv")
    per.sort(key=lambda x: (ORDER.index(x["form"]), x["orthogroup"]))
    for r in per:
        r["status"] = r["status"].replace("_", " ")
        r["genes"] = tx(r["genes"])
        for k in ("gene_ids", "same_gene_forms", "same_gene_orthogroups",
                  "same_gene_compartments"):
            r[k] = r[k].replace(",", ", ")
    blocks = [
        Block("(a) Summary by form",
              [("form", "Form", "text"), ("n_cloud_orthogroups", "Cloud orthogroups", "int"),
               ("with_outgroup_gene", f"With a {OG} gene", "int"),
               ("single_ingroup_gene", "With a single ingroup gene", "int"),
               ("same_gene_elsewhere", "Same gene in another form", "int"),
               ("same_gene_in_reference", "Same gene in quinquefasciatus", "int"),
               ("no_counterpart", "No counterpart", "int"),
               ("pct_same_gene_elsewhere", "Same gene in another form (%)", "pct1")], summ),
        Block("(b) Cloud orthogroups",
              [("orthogroup", "Orthogroup", "text"), ("form", "Form", "text"),
               ("n_ingroup_genes", "Ingroup genes", "int"),
               ("n_outgroup_genes", f"{OG} genes", "int"),
               ("genes", "Transcripts", "text"), ("gene_ids", "Genes", "text"),
               ("same_gene_forms", "Forms with the same gene", "text"),
               ("same_gene_orthogroups", "Orthogroups of the same gene", "text"),
               ("same_gene_compartments", "Compartments of those orthogroups", "text"),
               ("in_reference_elsewhere", "Same gene in quinquefasciatus", "bool"),
               ("status", "Status", "text")], per, autofilter=True),
    ]
    return ("S4 Cloud",
            "Supp. Table S4. Composition of the cloud orthogroups",
            f"Summary by form (a) and each of the {len(per):,} cloud orthogroups (b). Liftoff keeps the "
            "reference gene identifier on every transferred model, so models of the same gene "
            "can be traced across forms. 'Same gene in another form': another ingroup form "
            "carries a protein coding model with the same gene identifier, under a different "
            "orthogroup ('unassigned' when that model is in no orthogroup). 'Same gene in "
            "quinquefasciatus': that form is the reference. 'No counterpart': no model of the "
            "same gene in any other ingroup form.",
            "results/pangenome/cloud_composition.tsv; "
            "results/pangenome/cloud_composition_summary.tsv",
            blocks)


# ------------------------------------------------------------------ S5
def s5():
    summ = tsv("results/validation/absence_summary.tsv")
    calls = tsv("results/validation/absence_calls.tsv")
    order = {"cloud": 0, "shell": 1}
    calls.sort(key=lambda x: (order[x["compartment"]], x["orthogroup"], ALL5.index(x["absent_from"])))
    for r in calls:
        r["call"] = r["call"].replace("_", " ")
        r["match_basis"] = {"same_gene_id": "same gene identifier",
                            "overlap_identity": "near identical model"}.get(r["match_basis"], "")
        r["rep_gene"] = tx(r["rep_gene"])
        r["best_overlapping_gene"] = tx(r["best_overlapping_gene"])
        if r["call"] not in ("clustered elsewhere", "paralog only"):
            r["best_protein_identity"] = None
    blocks = [
        Block("(a) Summary by compartment",
              [("compartment", "Compartment", "text"), ("n_events", "Absence events", "int"),
               ("absent", "Absent", "int"), ("paralog_only", "Paralog only", "int"),
               ("clustered_elsewhere", "Clustered elsewhere", "int"),
               ("clustered_elsewhere_same_gene_id", "of which same gene identifier", "int"),
               ("clustered_elsewhere_overlap_identity", "of which near identical model", "int"),
               ("same_gene_id_unassigned", "Same gene identifier, target model in no orthogroup", "int"),
               ("same_gene_id_median_protein_identity",
                "Same gene identifier, median protein identity to the query", "f3"),
               ("unannotated_locus", "Unannotated locus", "int"), ("weak", "Weak", "int"),
               ("pct_supported", "Supported (%)", "pct1"), ("pct_artifact", "Artifact (%)", "pct1"),
               ("pct_weak", "Weak (%)", "pct1")], summ),
        Block("(b) Absence events",
              [("orthogroup", "Orthogroup", "text"), ("compartment", "Compartment", "text"),
               ("absent_from", "Form without the orthogroup", "text"),
               ("rep_gene", "Query transcript", "text"), ("rep_form", "Query form", "text"),
               ("prot_identity", "Protein probe identity", "f3"),
               ("prot_coverage", "Protein probe coverage", "f3"),
               ("dna_identity", "Locus probe identity", "f3"),
               ("dna_coverage", "Locus probe coverage", "f3"),
               ("call", "Call", "text"), ("match_basis", "Basis", "text"),
               ("best_overlapping_gene", "Matching model in the target", "text"),
               ("best_protein_identity", "Protein identity to the query", "f4"),
               ("other_orthogroup", "Orthogroup of the matching model", "text")],
              calls, autofilter=True),
    ]
    return ("S5 Absences",
            "Supp. Table S5. Validation of the gene absences implied by cloud and shell orthogroups",
            "Summary by compartment (a) and every absence event (b), an orthogroup and a form "
            "without it. The longest ingroup protein of the orthogroup (query) was aligned to the genome "
            f"of the form without it with miniprot {version('miniprot', '0.18')} (protein probe), and the genomic locus of "
            f"its gene, introns included, with minimap2 {version('minimap2', '2.31')} (locus probe); identity and coverage "
            "are for the best hit of each probe, 0 when there was none. Calls, in order of "
            "precedence: clustered elsewhere, basis same gene identifier (the target proteome holds "
            "a model with the query's gene identifier); absent (no hit with either probe); weak "
            "(hits, none with at least 0.80 identity and 0.70 protein or 0.50 locus coverage); and, "
            "for a hit passing those thresholds, unannotated locus (it overlaps no protein coding "
            "model), clustered elsewhere, basis near identical model (it overlaps a model with at "
            "least 0.90 protein identity to the query) or paralog only (it overlaps only less "
            "similar models). Protein identity is the number of identical residues over the length "
            "of the shorter protein after global alignment, or 0 when the shorter protein is less "
            "than 30% of the length of the longer. "
            "Supported absences are absent and paralog only events; artifacts are clustered "
            "elsewhere and unannotated locus events. Protein identity to the query is given for "
            "clustered elsewhere and paralog only calls.",
            "results/validation/absence_calls.tsv; results/validation/absence_summary.tsv",
            blocks)


# ------------------------------------------------------------------ S6
def s6():
    acc_label = {
        "sco_fastas": "Single copy orthologs (alignments before trimming)",
        "trimmed_alignments_kept": "Alignments kept after trimming (at least 50 columns)",
        "gene_trees_in_all_gene_trees": "Kept alignments with a gene tree",
        "trimmed_without_gene_tree": "Kept alignments without a gene tree",
        "trimmed_without_gene_tree_invariant": "  of which invariant",
        "trimmed_without_gene_tree_fewer_than_4_distinct": "  of which with fewer than four distinct sequences",
        "trimmed_invariant_total": "Kept alignments that are invariant",
        "trimmed_fewer_than_4_distinct_total": "Kept alignments with fewer than four distinct sequences",
    }
    acc = [{"m": acc_label[r["measure"]], "v": r["value"]}
           for r in tsv("results/phylo/locus_accounting.tsv")]
    topo = tsv("results/phylo/quartet_topology_counts.tsv")
    by_role = {r["role"]: dict(r) for r in topo}
    major, minor = by_role["major_discordant"], by_role["minor_discordant"]

    def with_label(name, r):
        return f"{name} ({r['iqtree_label']})" if r.get("iqtree_label") else name

    role = {"species_tree": "Species tree",
            "major_discordant": with_label("Major discordant", major),
            "minor_discordant": with_label("Minor discordant", minor)}
    n_gene_trees = sum(int(r["n_gene_trees"]) for r in topo)
    for r in topo:
        r["role"] = role[r["role"]]
    cf = [l for l in open(path("results/phylo/concord.cf.stat")) if not l.startswith("#")]
    cfd = dict(zip(cf[0].split(), cf[1].split()))
    cfrows = [
        {"m": "Gene trees informative for the branch (gN)", "v": cfd["gN"], "f": "int"},
        {"m": "Gene concordance factor, gCF (%)", "v": cfd["gCF"], "f": "pct2"},
        {"m": "Gene discordance factor 1, gDF1 (%)", "v": cfd["gDF1"], "f": "pct2"},
        {"m": "Gene discordance factor 2, gDF2 (%)", "v": cfd["gDF2"], "f": "pct2"},
        {"m": "Gene discordance factor, paraphyly, gDFP (%)", "v": cfd["gDFP"], "f": "pct2"},
        {"m": "Parsimony informative sites (sN)", "v": cfd["sN"], "f": "int"},
        {"m": "Site concordance factor, sCF (%)", "v": cfd["sCF"], "f": "pct2"},
        {"m": "Site discordance factor 1, sDF1 (%)", "v": cfd["sDF1"], "f": "pct2"},
        {"m": "Site discordance factor 2, sDF2 (%)", "v": cfd["sDF2"], "f": "pct2"},
        {"m": "UFBoot support", "v": cfd["Label"], "f": "int"},
        {"m": "Branch length, concatenated tree (substitutions per site)", "v": cfd["Length"], "f": "f4"},
    ]
    sup = tsv("results/phylo/quartet_asymmetry_by_support.tsv")
    sites = tsv("results/phylo/quartet_site_patterns.tsv")
    conc_row = next((r for r in sites if r["role"] == "concentration"), None)
    n_top = int(conc_row["n_loci"]) if conc_row else 0
    top_label = f"Top 1% of loci by informative sites ({n_top:,} loci)"
    robust = tsv("results/phylo/quartet_robustness.tsv")
    for r in robust:
        r["subset"] = (r["subset"][0].upper() + r["subset"][1:]).replace(">=", "\u2265")
    label = {
        "species_tree": ("All loci, summed sites", True),
        "major_discordant_gene_trees": ("All loci, summed sites", True),
        "minor_discordant_gene_trees": ("All loci, summed sites", True),
        "locus_majority_species_tree": ("Locus majority votes", True),
        "locus_majority_major_discordant_gene_trees": ("Locus majority votes", True),
        "locus_majority_minor_discordant_gene_trees": ("Locus majority votes", True),
        "top_loci_species_tree": (top_label, True),
        "top_loci_major_discordant_gene_trees": (top_label, True),
        "top_loci_minor_discordant_gene_trees": (top_label, True),
    }
    test_label = {
        "discordant sites: gene tree major vs minor":
            "Test: summed discordant sites, major vs minor split, all loci",
        "locus majority: gene tree major vs minor":
            "Test: locus majority votes, major vs minor split",
        "discordant sites outside the top 1% of loci: gene tree major vs minor":
            "Test: summed discordant sites, major vs minor split, without the top 1% of loci",
    }
    srows = []
    for r in sites:
        d = {"n_sites": r["n_informative_sites"], "pct": r["pct"], "p": r["binomial_p"],
             "lo": r["ci95_low"], "hi": r["ci95_high"], "loci": r["n_loci"],
             "loci0": r["n_loci_no_informative_sites"], "split": ""}
        if r["role"] in label:
            d["measure"] = label[r["role"]][0]
            d["split"] = forms(r["split"]).replace("+", " + ")
        elif r["role"] == "binomial_test":
            d["measure"] = test_label[r["split"]]
        elif r["role"] == "locus_majority_undecided":
            d["measure"] = "Loci without a majority split (ties or no informative sites)"
        elif r["role"] == "median":
            d["measure"] = "Informative sites per locus, median"
            d["pct"] = None
        elif r["role"] == "concentration":
            d["measure"] = "Top 1% of loci by informative sites: share of all informative sites"
        elif r["role"] == "loci_with_a_model_not_intact":
            d["measure"] = (f"{r['split'][0].upper()}{r['split'][1:]} ({int(r['n_loci']):,} loci): "
                            "loci with a gene model that is not intact")
            d["loci"], d["n_sites"] = r["n_loci_not_intact"], None
        else:
            raise SystemExit(f"unknown role {r['role']}")
        srows.append(d)
    for r in topo:
        r["split"] = forms(r["split"]).replace("+", " + ")
    for r in sup:
        r["major_split"] = forms(r["major_split"]).replace("+", " + ")
        r["minor_split"] = forms(r["minor_split"]).replace("+", " + ")
    blocks = [
        Block("(a) Single copy loci and gene trees",
              [("m", "Measure", "text"), ("v", "Loci", "int")], acc),
        Block("(b) Gene trees supporting each unrooted topology",
              [("split", "Split", "text"), ("role", "Topology", "text"),
               ("n_gene_trees", "Gene trees", "int"), ("pct", "Gene trees (%)", "pct2"),
               ("iqtree_label", "IQ-TREE label", "text")], topo),
        Block(f"(c) Concordance factors of the internal branch (IQ-TREE {version('iqtree', '3.1.3')})",
              [("m", "Measure", "text"), ("v", "Value", "gen")], cfrows),
        Block("(d) Discordant gene trees by minimum UFBoot support of the internal branch",
              [("min_ufboot", "Minimum UFBoot", "int"), ("n_gene_trees", "Gene trees", "int"),
               ("n_concordant", "Concordant", "int"), ("pct_concordant", "Concordant (%)", "pct2"),
               ("n_major", "Major discordant", "int"), ("n_minor", "Minor discordant", "int"),
               ("major_share_of_discordant", "Major share of discordant", "f4"),
               ("binomial_p", "Binomial P", "p"), ("ci95_low", "95% CI, low", "f4"),
               ("ci95_high", "95% CI, high", "f4")], sup),
        Block("(e) Parsimony informative sites and locus votes",
              [("measure", "Measure", "text"), ("split", "Split", "text"),
               ("n_sites", "Informative sites", "int"), ("loci", "Loci", "int"),
               ("pct", "Share (%)", "pct2"), ("p", "Binomial P", "p"),
               ("lo", "95% CI, low", "f4"), ("hi", "95% CI, high", "f4"),
               ("loci0", "Loci without informative sites", "int")], srows),
        Block("(f) Gene tree test on resolved gene trees and on loci with intact gene models",
              [("subset", "Gene trees", "text"), ("n_gene_trees", "Gene trees (n)", "int"),
               ("n_concordant", "Concordant", "int"), ("pct_concordant", "Concordant (%)", "pct2"),
               ("n_major", "Major discordant", "int"), ("n_minor", "Minor discordant", "int"),
               ("major_share_of_discordant", "Major share of discordant", "f4"),
               ("binomial_p", "Binomial P", "p"), ("ci95_low", "95% CI, low", "f4"),
               ("ci95_high", "95% CI, high", "f4")], robust),
    ]
    # the concordance factor block mixes formats; format each value by its row
    blocks[2].row_formats = [r["f"] for r in cfrows]

    def split_text(r):
        return forms(r["split"]).replace("+", " + ")

    def label_text(r):
        return f", is {r['iqtree_label']} in IQ-TREE's labels" if r.get("iqtree_label") else ""
    return ("S6 Quartet tests",
            "Supp. Table S6. Single copy loci and tests of the four taxon quartet",
            "(a) Single copy loci from orthogroups to gene trees; distinct sequences are counted "
            "on the trimmed alignments. The remaining parts test the symmetry of the two "
            "discordant topologies of the unrooted quartet. Splits are written as the two pairs of "
            "forms separated by a vertical bar; the major discordant split, "
            f"{split_text(major)}{label_text(major)}, and the minor discordant split, "
            f"{split_text(minor)}{label_text(minor)}. (b) Gene trees of the {n_gene_trees:,} "
            "loci with a gene tree. (c) Concordance factors of the internal "
            "branch. (d) The same counts among gene trees whose internal branch reached each "
            "UFBoot value; the major share is tested against 0.5 with an exact binomial test, with "
            "Clopper Pearson 95% confidence intervals. (e) "
            "Parsimony informative sites summed over loci; locus majority votes, in which each "
            "locus counts once for the split supported by most of its informative sites (ties and "
            f"loci without informative sites set aside); and the {n_top:,} loci (1%) with the most "
            "informative sites. Share (%) is the share of sites or loci in the group that support "
            "the split; for tests, the share of discordant sites or votes supporting the major "
            "discordant split, and the confidence interval is for that share (0 to 1 scale). "
            "(f) The gene tree test on all gene trees; on gene trees whose internal branch is "
            "longer than IQ-TREE's minimum length (resolved), and on the rest, which have no "
            "substitution supporting any resolution; on gene trees without and with a terminal "
            "branch longer than 0.1, the trees most exposed to long branch attraction; and on "
            "loci whose four gene models are "
            "intact (reference model neither partial nor with a RefSeq exception, and a valid ORF "
            "for each transferred model), and on the rest. Part (e) also gives how many of the "
            "top 1% of loci, and of the other loci, hold a gene model that is not intact.",
            "results/phylo/locus_accounting.tsv; results/phylo/quartet_topology_counts.tsv; "
            "results/phylo/concord.cf.stat; results/phylo/quartet_asymmetry_by_support.tsv; "
            "results/phylo/quartet_site_patterns.tsv",
            blocks)


# ------------------------------------------------------------------ S7
def s7():
    bl = tsv("results/phylo/branch_length_summary.tsv")
    for r in bl:
        r["branch"] = "internal" if r["branch"] == "internal" else FORM[r["branch"]]
    ident = tsv("results/phylo/sco_pairwise_identity.tsv")
    ani = tsv("results/synteny/ani_pairs.tsv")
    syn = {frozenset([r["sample1"], r["sample2"]]): r for r in tsv("results/synteny/synteny_summary.tsv")}
    wg = []
    for r in ani:
        s = syn.get(frozenset([r["sample1"], r["sample2"]]))
        wg.append({"pair": f"{FORM[r['sample1']]} vs {FORM[r['sample2']]}",
                   "ani": r["ani"] if r["ani"] != "NA" else "no estimate",
                   "afq": r["align_fraction_query"], "afr": r["align_fraction_ref"],
                   "mm": s["mean_alignment_identity_pct"] if s else None,
                   "bp": s["total_aligned_bp"] if s else None})
    for r in ident:
        r["taxon_a"], r["taxon_b"] = FORM[r["taxon_a"]], FORM[r["taxon_b"]]
    na = [r for r in ani if r["ani"] in ("NA", "")]
    og_pairs = [r for r in ani if OUTGROUP in (r["sample1"], r["sample2"])]
    if na and len(na) == len(og_pairs) and all(OUTGROUP in (r["sample1"], r["sample2"]) for r in na):
        ani_note = f" skani returned no estimate for pairs with {OG}."
    elif na:
        ani_note = (" skani returned no estimate for "
                    + ", ".join(f"{FORM[r['sample1']]} vs {FORM[r['sample2']]}" for r in na) + ".")
    else:
        ani_note = ""
    blocks = [
        Block("(a) Branch lengths, concatenated tree and per locus gene trees (substitutions per site)",
              [("branch", "Branch", "text"), ("concat_tree", "Concatenated tree", "f5"),
               ("n_gene_trees", "Gene trees", "int"), ("gene_tree_median", "Gene trees, median", "f5"),
               ("gene_tree_mean", "Gene trees, mean", "f5"),
               ("gene_tree_p90", "Gene trees, 90th percentile", "f5"),
               ("n_above_0.1", "Gene trees with the branch above 0.1", "int"),
               ("share_above_0.1", "Share of gene trees above 0.1", "f5"),
               ("n_minimum", "Gene trees with the branch at the minimum length", "int"),
               ("share_minimum", "Share of gene trees at the minimum length", "f5")], bl),
        Block("(b) Pairwise protein identity over the trimmed single copy alignments",
              [("taxon_a", "Form", "text"), ("taxon_b", "Form", "text"),
               ("n_loci", "Loci", "int"), ("median_identity", "Median identity", "f5"),
               ("mean_identity", "Mean identity", "f5"),
               ("p05_identity", "5th percentile", "f5"),
               ("share_below_0.95", "Share of loci below 0.95", "f4")], ident),
        Block("(c) Whole genome identity",
              [("pair", "Pair", "text"), ("ani", "skani ANI (%)", "pct2"),
               ("afq", "skani aligned fraction, query (%)", "pct2"),
               ("afr", "skani aligned fraction, reference (%)", "pct2"),
               ("mm", "minimap2 alignment identity (%)", "pct2"),
               ("bp", "minimap2 aligned length (bp)", "int")], wg),
    ]
    return ("S7 Divergence",
            "Supp. Table S7. Divergence between the forms",
            "(a) Branch lengths of the unrooted concatenated tree and of the per locus gene trees; "
            "a branch at IQ-TREE's minimum length (10^-6) has no inferred substitution. (b) "
            "Pairwise protein identity for each pair of forms, computed for each trimmed single "
            "copy alignment over columns where both forms have a residue. (c) Whole genome "
            "identity: skani (skani dist -s 80, first assembly of the pair given first) on the "
            "complete assemblies, with the aligned fractions as skani labels them; minimap2 "
            f"{version('minimap2', '2.31')} (-x asm20) identity is the length weighted mean of one "
            "minus the divergence of each alignment between the three chromosome scale sequences "
            f"of the two assemblies.{ani_note}",
            "results/phylo/branch_length_summary.tsv; results/phylo/sco_pairwise_identity.tsv; "
            "results/synteny/ani_pairs.tsv; results/synteny/synteny_summary.tsv",
            blocks)


# ------------------------------------------------------------------ S8
def s8():
    res = {}
    for line in open(path("results/cafe/output/Gamma_results.txt")):
        if ":" in line:
            k, v = line.split(":", 1)
            res[k.strip()] = v.strip()
    tree = open(path("results/cafe/ultrametric_tree.nwk")).read().strip()
    tb = {r["metric"]: r["value"] for r in tsv("results/cafe/transfer_bias_summary.tsv")}
    model = [
        {"m": "Model", "v": "gamma, three rate categories"},
        {"m": "Tree (branch lengths assumed)", "v": forms(tree)},
        {"m": "Final likelihood (-lnL)", "v": res["Model Gamma Final Likelihood (-lnL)"]},
        {"m": "Birth and death rate (lambda)", "v": float(res["Lambda"])},
        {"m": "Gamma shape (alpha)", "v": float(res["Alpha"])},
        {"m": "Families in the gene count table", "v": tb["n_families_in_counts"]},
        {"m": "Families tested (present at the root)", "v": tb["n_families"]},
        {"m": "Significant families (P < 0.05)", "v": tb["n_significant"]},
    ]
    br = tsv("results/cafe/branch_summary.tsv")
    for r in br:
        r["taxon"] = forms(r["taxon"]).replace("+", " + ")
        r["node_label"] = forms(r["node_label"])
    bycn = tsv("results/cafe/transfer_bias_by_copy_number.tsv")
    for r in bycn:
        if r["ref_copies"] == "11+":
            r["ref_copies"] = "more than 10"
        r["ref_copies"] = r["ref_copies"].replace("-", " to ")
    lin = tsv("results/cafe/transfer_bias_by_lineage.tsv")
    for r in lin:
        r["taxon"] = FORM[r["taxon"]]
    labels = [
        ("n_families_single_copy", "Families with one reference copy", "int"),
        ("n_families_multi_copy", "Families with two or more reference copies", "int"),
        ("n_families_zero_ref_copy", "Families with no reference copy", "int"),
        ("n_significant_single_copy", "Significant, one reference copy", "int"),
        ("n_significant_multi_copy", "Significant, two or more reference copies", "int"),
        ("n_significant_zero_ref_copy", "Significant, no reference copy", "int"),
        ("pct_significant_single_copy", "Significant (%), one reference copy", "pct3"),
        ("pct_significant_multi_copy", "Significant (%), two or more reference copies", "pct3"),
        ("pct_significant_zero_ref_copy", "Significant (%), no reference copy", "pct3"),
        ("retention_single_copy", "Mean retention, one reference copy", "f4"),
        ("retention_multi_copy", "Mean retention, two or more reference copies", "f4"),
        ("odds_ratio_multi_vs_single_copy", "Odds ratio, significant, two or more against one copy", "f3"),
        ("risk_ratio_multi_vs_single_copy", "Relative risk, significant, two or more against one copy", "f3"),
        ("fisher_p_two_sided", "Fisher's exact test P (two sided)", "p"),
        ("mean_deficit_significant", "Mean copy deficit, significant families", "f4"),
        ("mean_deficit_not_significant", "Mean copy deficit, other families", "f4"),
        ("mannwhitney_p_deficit_two_sided", "Mann Whitney P, copy deficit (two sided)", "p"),
        ("mannwhitney_p_ref_copies_two_sided", "Mann Whitney P, reference copies (two sided)", "text"),
        ("per_lineage_pearson_r", "Pearson r, net contraction against total copy deficit (4 lineages)", "f4"),
        ("per_lineage_p", "P for that correlation", "f5"),
    ]
    tests = [{"m": lab, "v": tb[k], "f": f} for k, lab, f in labels]
    counts = {r["Family ID"]: r for r in tsv("results/cafe/gene_counts_filtered.tsv")}
    sig = []
    for r in tsv("results/cafe/significant_families.tsv"):
        c = counts[r["orthogroup"]]
        d = {"og": r["orthogroup"], "p": r["pvalue"]}
        for s in ORDER:
            d[s] = c[s]
        sig.append(d)
    sig.sort(key=lambda x: x["og"])
    blocks = [
        Block("(a) Model", [("m", "Measure", "text"), ("v", "Value", "gen")], model),
        Block("(b) Changes per branch",
              [("taxon", "Branch", "text"), ("node_label", "CAFE node", "text"),
               ("increase", "Increases", "int"), ("decrease", "Decreases", "int"),
               ("total", "Total", "int"), ("ratio_inc_dec", "Increases per decrease", "f4"),
               ("binomial_p", "Binomial P, increases against decreases", "p")], br),
        Block("(c) Families by number of copies in the reference",
              [("ref_copies", "Reference copies", "text"), ("n_families", "Families", "int"),
               ("mean_lifted_copies", "Mean copies per transferred genome", "f4"),
               ("mean_retention", "Mean retention", "f4"), ("mean_deficit", "Mean copy deficit", "f4"),
               ("n_significant", "Significant", "int"), ("pct_significant", "Significant (%)", "pct3")],
              bycn),
        Block("(d) Copies and CAFE changes per lineage",
              [("taxon", "Form", "text"), ("total_copies", "Total copies", "int"),
               ("copy_deficit_vs_reference", "Copy deficit against the reference", "int"),
               ("increase", "Increases", "int"), ("decrease", "Decreases", "int"),
               ("net_contraction", "Net contraction", "int")], lin),
        Block("(e) Tests", [("m", "Measure", "text"), ("v", "Value", "gen")], tests),
        Block("(f) Significant families",
              [("og", "Orthogroup", "text"), ("p", "CAFE P", "f3")] +
              [(s, f"{FORM[s]} copies", "int") for s in ORDER], sig, autofilter=True),
    ]
    blocks[4].row_formats = [t["f"] for t in tests]
    blocks[0].row_formats = ["text", "text", "gen", "f5", "f5", "int", "int", "int"]
    n_in = int(float(tb["n_families_in_counts"]))
    n_tested = int(float(tb["n_families"]))
    rooted = os.path.exists(path("results/phylo/rooted/rooted_tree.treefile"))
    tree_note = ("the tree is the species tree rooted with the outgroup, and branch lengths were "
                 "assumed because divergence times are not known" if rooted else
                 "branch lengths were assumed because neither the root nor divergence times are "
                 "known")
    return ("S8 CAFE",
            "Supp. Table S8. CAFE gene family analysis and the annotation transfer control",
            f"CAFE 5 on the counts of genes of the four ingroup forms in {n_in:,} orthogroups; "
            "CAFE tests only families present at the root and set aside the "
            f"{n_in - n_tested:,} without genes on one side of it, so the analyses below cover "
            f"the {n_tested:,} tested families. (a) Model and tree; {tree_note}. (b) Families "
            "increasing and decreasing on each branch, with a "
            "binomial test of increases against decreases. (c) Families binned by the number of "
            "copies in the Cx. quinquefasciatus reference: mean copies per transferred genome "
            "(pallens, molestus and pipiens), mean share of reference copies retained (retention), "
            "mean copies lost (deficit) and the share of families CAFE called significant. (d) "
            "Total copies of each form over the tested families and its deficit relative to the "
            "reference, with CAFE's increases and decreases on its terminal branch. (e) Tests "
            f"comparing families with one and with two or more reference copies. (f) The {len(sig):,} "
            "families with P < 0.05 and their copies per form; CAFE reports P to three decimals.",
            "results/cafe/output/Gamma_results.txt; results/cafe/ultrametric_tree.nwk; "
            "results/cafe/branch_summary.tsv; results/cafe/transfer_bias_by_copy_number.tsv; "
            "results/cafe/transfer_bias_by_lineage.tsv; results/cafe/transfer_bias_summary.tsv; "
            "results/cafe/significant_families.tsv; results/cafe/gene_counts_filtered.tsv",
            blocks)


# ------------------------------------------------------------------ S9
def s9():
    syn = tsv("results/synteny/synteny_summary.tsv")
    for r in syn:
        r["pair"] = pair_name(r["pair"])
        r["top_inversions"] = r["top_inversions"].replace("; ", "; ")
    inv = tsv("results/synteny/inversions_detail.tsv")
    chrom = {}
    for r in inv:
        chrom[(r["sample1"], r["seqid"])] = r["chromosome"]
    ori = tsv("results/synteny/chromosome_orientation.tsv")
    for r in ori:        # name the second assembly's sequences through the first
        if (r["sample1"], r["seqid1"]) in chrom:
            chrom.setdefault((r["sample2"], r["seqid2"]), chrom[(r["sample1"], r["seqid1"])])
    for r in ori:
        r["chromosome"] = chrom.get((r["sample1"], r["seqid1"]), "")
        r["sample1"], r["sample2"] = FORM[r["sample1"]], FORM[r["sample2"]]
    for r in inv:
        r["sample1"], r["sample2"] = FORM[r["sample1"]], FORM[r["sample2"]]
    rec = tsv("results/synteny/inversion_recurrence.tsv")
    for r in rec:
        g = r["genome_in_every_pair"]
        r["common"] = FORM.get(g, g)
        r["pairs"] = "; ".join(pair_name(p) for p in r["pairs"].split(";"))
        r["members"] = "; ".join(pair_name(m.split(":", 1)[0]) + ": " + m.split(":", 1)[1]
                                 for m in r["members"].split(";"))
        r["any_at_least_large_span"] = "yes" if r["any_at_least_large_span"] == "True" else "no"
    blocks = [
        Block("(a) Pairs of assemblies",
              [("pair", "Pair", "text"),
               ("n_shared_genes_top_sequences", "Shared genes on the chromosome scale sequences", "int"),
               ("n_shared_genes_chr1_3", "Anchors on homologous chromosomes", "int"),
               ("n_shared_genes_other_chromosome", "Shared genes on another chromosome", "int"),
               ("pct_shared_genes_other_chromosome", "Shared genes on another chromosome (%)", "pct2"),
               ("pct_collinear_anchors", "Collinear anchors (%)", "pct2"),
               ("n_chrom_pairs_reverse_oriented", "Chromosome pairs in opposite orientation", "int"),
               ("n_inversions_intra_chr", "Inversions of at least 100 kb", "int"),
               ("max_inversion_size_bp", "Largest inversion (bp)", "int"),
               ("top_inversions", "Three largest inversions (first assembly coordinates)", "text")],
              syn),
        Block("(b) Chromosome pairing and orientation",
              [("sample1", "First assembly", "text"), ("sample2", "Second assembly", "text"),
               ("chromosome", "Chromosome", "text"), ("seqid1", "Sequence, first", "text"),
               ("seqid2", "Sequence, second", "text"), ("n_anchors", "Anchors", "int"),
               ("lis_forward", "Longest increasing run", "int"),
               ("lis_reverse", "Longest decreasing run", "int"),
               ("orientation", "Relative orientation", "text")], ori),
        Block("(c) Inversions of at least 100 kb",
              [("sample1", "First assembly", "text"), ("sample2", "Second assembly", "text"),
               ("seqid", "Sequence, first", "text"), ("chromosome", "Chromosome", "text"),
               ("start_bp", "Start (bp)", "int"), ("end_bp", "End (bp)", "int"),
               ("span_bp", "Span (bp)", "int"), ("n_genes", "Anchor genes", "int"),
               ("rel_start", "Relative start", "f3"), ("rel_end", "Relative end", "f3"),
               ("first_gene", "First gene", "text"), ("last_gene", "Last gene", "text"),
               ("recurrence_cluster", "Recurrence cluster", "int"),
               ("n_pairs_in_cluster", "Pairs in cluster", "int")], inv, autofilter=True),
        Block("(d) Recurrence clusters",
              [("cluster", "Cluster", "int"), ("n_inversions", "Inversions", "int"),
               ("n_pairs", "Pairs", "int"), ("common", "Assembly in every pair", "text"),
               ("chromosomes", "Chromosome", "text"), ("max_span_bp", "Largest span (bp)", "int"),
               ("any_at_least_large_span", "Any of at least 1 Mb", "text"),
               ("pairs", "Pairs of assemblies", "text"), ("members", "Inversions (first assembly coordinates)", "text")],
              rec),
    ]
    return ("S9 Synteny",
            "Supp. Table S9. Synteny between the ingroup assemblies",
            "Synteny from the positions (gene midpoints) of genes, coding and non coding, shared by two "
            "assemblies on their three chromosome scale sequences (reference positions in "
            "Cx. quinquefasciatus, transferred positions in the other forms). (a) Summary per pair: anchors "
            "are shared genes on homologous chromosomes, and collinear anchors lie on the longest "
            "increasing run after orientation. (b) Chromosome pairing and the relative orientation "
            "of each pair of homologous sequences, chosen as the direction with the longer monotone "
            "run of anchors. (c) Inversions: runs of at least three anchors in decreasing order "
            "spanning at least 100 kb, with coordinates of the first and last anchor genes on the "
            "sequence of the first assembly and their relative positions along it (0 to 1). (d) "
            "Recurrence clusters: inversions from different pairs joined by single linkage when their "
            "anchor gene sets have a Jaccard index of at least 0.5; for clusters found in two or more pairs, "
            "the assembly present in every pair, or none. A cluster found in all three pairs that "
            "include one assembly is the pattern expected for a rearrangement in, or an assembly "
            "error of, that assembly. Chromosomes are numbered after "
            "Cx. quinquefasciatus.",
            "results/synteny/synteny_summary.tsv; results/synteny/chromosome_orientation.tsv; "
            "results/synteny/inversions_detail.tsv; results/synteny/inversion_recurrence.tsv",
            blocks)


# ------------------------------------------------------------------ S10
def s10():
    rows = []
    pat = {
        "masked": r"bases masked:\s+(\d+) bp\s+\(\s*([\d.]+) %\)",
        "retro": r"^Retroelements\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "line": r"^\s+LINEs:\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "ltr": r"^\s+LTR elements:\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "dna": r"^DNA transposons\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "rc": r"^Rolling-circles\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "uncl": r"^Unclassified:\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "inter": r"^Total interspersed repeats:\s+\d+ bp\s+([\d.]+) %",
        "srna": r"^Small RNA:\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "sat": r"^Satellites:\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "simple": r"^Simple repeats:\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "low": r"^Low complexity:\s+\d+\s+\d+ bp\s+([\d.]+) %",
        "length": r"total length:\s+(\d+) bp",
    }
    for s in ORDER:
        text = open(path(f"results/repeats/{s}/{s}.fasta.tbl")).read()
        d = {"form": FORM[s]}
        for k, p in pat.items():
            m = re.search(p, text, re.M)
            if not m:
                raise SystemExit(f"{s}: no match for {k}")
            if k == "masked":
                d["masked_bp"], d["masked"] = m.group(1), m.group(2)
            else:
                d[k] = m.group(1)
        rows.append(d)
    cols = [("form", "Form", "text"), ("length", "Total length (bp)", "int"),
            ("masked_bp", "Bases masked (bp)", "int"), ("masked", "Masked (%)", "pct2"),
            ("inter", "Interspersed repeats (%)", "pct2"), ("retro", "Retroelements (%)", "pct2"),
            ("line", "LINEs (%)", "pct2"), ("ltr", "LTR elements (%)", "pct2"),
            ("dna", "DNA transposons (%)", "pct2"), ("rc", "Rolling circles (%)", "pct2"),
            ("uncl", "Unclassified (%)", "pct2"), ("srna", "Small RNA (%)", "pct2"),
            ("sat", "Satellites (%)", "pct2"), ("simple", "Simple repeats (%)", "pct2"),
            ("low", "Low complexity (%)", "pct2")]
    return ("S10 Repeats",
            "Supp. Table S10. Repeat content of the ingroup assemblies",
            f"RepeatMasker {version('repeatmasker', '4.1.7').split('-')[0]} summaries for the four ingroup assemblies, each masked with one "
            f"RepeatModeler {version('repeatmodeler', '2.0.7')} library built from the Cx. quinquefasciatus assembly. Percentages "
            "are of the total assembly length. "
            "Interspersed repeats comprise retroelements, DNA transposons, rolling circles and "
            "unclassified repeats; most interspersed repeats in the de novo library were "
            f"unclassified. No SINEs or Penelope elements were found. {OG} was not masked.",
            "results/repeats/<form>/<form>.fasta.tbl",
            [Block(None, cols, rows)])


# ------------------------------------------------------------------ S11
def s11():
    names = {"P450": "Cytochrome P450s", "OR": "Odorant receptors",
             "OBP": "Odorant binding proteins", "CCE": "Carboxylesterases",
             "GR": "Gustatory receptors", "IR": "Ionotropic receptors",
             "immune": "Immune genes", "GST": "Glutathione S transferases",
             "CSP": "Chemosensory proteins"}
    wide = tsv("results/functional/key_families_wide.tsv")
    wide.sort(key=lambda r: -int(r["Cx_quinquefasciatus"]))
    fam_order = [r["family"] for r in wide]
    for r in wide:
        r["name"] = names.get(r["family"], r["family"])
    long = tsv("results/functional/key_families_long.tsv")
    per = collections.OrderedDict()
    for r in long:
        d = per.setdefault(r["orthogroup"], {"og": r["orthogroup"], "family": r["family"],
                                             **{s: 0 for s in ORDER}})
        d[r["form"]] = int(r["copies"])
    rows = sorted(per.values(), key=lambda d: (fam_order.index(d["family"]), d["og"]))
    for d in rows:
        d["name"] = names.get(d["family"], d["family"])
    blocks = [
        Block("(a) Genes per family",
              [("name", "Family", "text")] + [(s, FORM[s], "int") for s in ORDER], wide),
        Block("(b) Orthogroups assigned to each family",
              [("og", "Orthogroup", "text"), ("name", "Family", "text")] +
              [(s, FORM[s], "int") for s in ORDER], rows, autofilter=True),
    ]
    return ("S11 Key families",
            "Supp. Table S11. Key gene family assignments",
            "Orthogroups assigned to nine gene families from eggNOG mapper 2.1.12 annotations "
            "(eggNOG 5.0) by weighted keyword patterns on descriptions and gene names; each "
            "orthogroup takes the family with the most annotated members, provided that family "
            "accounts for at least two genes or a quarter of the orthogroup's ingroup genes. Immune "
            "genes are Toll like receptors, peptidoglycan recognition proteins, thioester containing "
            "proteins, IMD and antimicrobial peptides. "
            "(a) Genes per form in the assigned orthogroups (as in Table 4). (b) Genes of each "
            "form in each assigned orthogroup. Counts in the three transferred gene sets are "
            "lower bounds, because copies are lost during transfer.",
            "results/functional/key_families_wide.tsv; results/functional/key_families_long.tsv",
            blocks)


# ------------------------------------------------------------------ S12
def s12():
    """The outgroup analyses (V5): five taxon loci, the rooted species tree,
    rooted gene tree topologies and the D statistics. None when absent."""
    need = ["results/phylo/rooted/rooted_summary.tsv",
            "results/phylo/rooted/rooted_topology_counts.tsv",
            "results/phylo/rooted/sco5_accounting.tsv",
            "results/phylo/rooted/alignment_summary.tsv",
            "results/phylo/dstat/d_statistics.tsv"]
    missing = [p for p in need if not os.path.exists(path(p))]
    if missing:
        print(f"S12 skipped, missing: {', '.join(missing)}")
        return None
    item_label = {
        "single_copy_in_all_taxa": "Orthogroups with one gene in each of the five taxa",
        "removed_no_protein": "Removed: a protein missing from the proteome",
        "removed_no_cds": "Removed: a coding sequence missing",
        "removed_cds_mismatch": "Removed: a coding sequence whose translation matches its "
                                "protein at less than 95% identity in every frame",
        "written": "Loci with proteins and coding sequences",
        "alignments": "Protein alignments (MAFFT)",
        "no_columns_kept": "Codon alignment not made: trimAl kept no column",
        "alignment_not_translation": "Codon alignment not made: alignment differs from the translation",
        "cds_length_mismatch": "Codon alignment not made: coding sequence length differs",
        "missing_files": "Codon alignment not made: files missing",
        "trimmed_protein_alignments_at_least_50_columns":
            "Trimmed protein alignments of at least 50 columns (rooted tree)",
    }
    acc = []
    for r in tsv("results/phylo/rooted/sco5_accounting.tsv"):
        if r["sample"]:
            if r["item"] == "cds_translation_frame_not_0" and r["value"] != "0":
                acc.append({"m": f"Coding sequences read from the second or third base "
                                 f"(partial models), {FORM[r['sample']]}", "v": r["value"]})
            continue
        acc.append({"m": item_label.get(r["item"], r["item"]), "v": r["value"]})
    for r in tsv("results/phylo/rooted/alignment_summary.tsv"):
        label = item_label.get(r["item"], r["item"])
        if r["item"] == "written":
            label = "Codon alignments"
        acc.append({"m": label, "v": r["value"]})
    summ = {r["item"]: r["value"] for r in tsv("results/phylo/rooted/rooted_summary.tsv")}
    tree_rows = [
        {"m": "Loci in the concatenated alignment", "v": summ.get("concatenated alignment loci")},
        {"m": "Columns in the concatenated alignment",
         "v": summ.get("concatenated alignment columns")},
        {"m": "Rooted topology of the ingroup", "v": forms(summ.get("rooted ingroup topology"))},
        {"m": "Root of the ingroup", "v": forms(summ.get("root position"))},
    ]
    clades = []
    for k, v in summ.items():
        if k.startswith("(") and k.endswith(" UFBoot"):
            name = k[:-len(" UFBoot")]
            row = {"clade": forms(name).replace(",", ", "), "ufboot": v,
                   "length": summ.get(f"{name} branch length")}
            for f in ("gCF", "gDF1", "gDF2", "gDFP", "gN", "sCF", "sDF1", "sDF2", "sN"):
                row[f] = summ.get(f"{name} {f}")
            clades.append(row)
    topo = tsv("results/phylo/rooted/rooted_topology_counts.tsv")
    for r in topo:
        r["rooted_topology"] = forms(r["rooted_topology"]).replace(",", ", ")
        r["quartet_split"] = forms(r["quartet_split"]).replace("+", " + ")
    dstat = tsv("results/phylo/dstat/d_statistics.tsv")
    set_label = {"all_loci": "all loci", "intact_loci": "loci with four intact models"}
    class_label = {"all_sites": "all sites", "third_positions": "third codon positions"}
    for r in dstat:
        r["statistic"] = forms(r["statistic"]).replace(",", ", ").replace(";", "; ")
        r["locus_set"] = set_label.get(r["locus_set"], r["locus_set"])
        r["site_class"] = class_label.get(r["site_class"], r["site_class"])
        r["excess_derived_sharing"] = (forms(r["excess_derived_sharing"]).replace("+", " + ")
                                       if r.get("abs_Z_at_least_3") == "True" else "none")
    blocks = [
        Block("(a) Five taxon single copy loci", [("m", "Measure", "text"), ("v", "Value", "int")],
              acc),
        Block("(b) Species tree rooted with the outgroup",
              [("m", "Measure", "text"), ("v", "Value", "gen")], tree_rows),
        Block("(c) Internal branches of the rooted tree",
              [("clade", "Clade", "text"), ("ufboot", "UFBoot", "int"),
               ("gCF", "gCF (%)", "pct2"), ("gDF1", "gDF1 (%)", "pct2"),
               ("gDF2", "gDF2 (%)", "pct2"), ("gDFP", "gDFP (%)", "pct2"),
               ("gN", "Gene trees (gN)", "int"), ("sCF", "sCF (%)", "pct2"),
               ("sDF1", "sDF1 (%)", "pct2"), ("sDF2", "sDF2 (%)", "pct2"),
               ("sN", "Sites (sN)", "pct2"), ("length", "Branch length", "f5")], clades),
        Block("(d) Rooted topologies of the five taxon gene trees",
              [("rooted_topology", "Rooted ingroup topology", "text"),
               ("quartet_split", "Unrooted ingroup split", "text"),
               ("is_species_tree", "Species tree", "bool"),
               ("n_all", "Gene trees", "int"), ("pct_all", "Gene trees (%)", "pct2"),
               ("n_resolved", "Resolved gene trees", "int"),
               ("pct_resolved", "Resolved gene trees (%)", "pct2"),
               ("n_ufboot95", "Gene trees, UFBoot \u2265 95", "int"),
               ("pct_ufboot95", "Gene trees, UFBoot \u2265 95 (%)", "pct2")], topo),
        Block("(e) D statistics",
              [("test", "Test", "text"), ("statistic", "Statistic", "text"),
               ("locus_set", "Loci", "text"), ("site_class", "Sites", "text"),
               ("consistent_with_rooted_tree", "P1 and P2 sisters in the rooted tree", "bool"),
               ("n_loci", "Loci (n)", "int"), ("n_blocks", "Jackknife blocks", "int"),
               ("n_sites", "Sites (n)", "int"), ("BBAA", "BBAA", "int"),
               ("ABBA", "ABBA", "int"), ("BABA", "BABA", "int"), ("D", "D", "f4"),
               ("se", "SE", "f4"), ("Z", "Z", "f3"), ("p", "P", "p"),
               ("p_bonferroni", "P, Bonferroni", "p"),
               ("excess_derived_sharing", "Pair sharing more derived alleles (|Z| \u2265 3)",
                "text")], dstat),
    ]
    blocks[1].row_formats = ["int", "int", "text", "text"]
    block_mb = VALUES.get(("parameters", "D statistics", ""), "")
    m = re.search(r"over (\d+) Mb windows", block_mb)
    window = f"{m.group(1)} Mb" if m else "5 Mb"
    return ("S12 Outgroup tests",
            "Supp. Table S12. Rooted species tree and D statistics with the outgroup",
            f"Analyses of the orthogroups with one gene in each ingroup form and in {OG}. (a) "
            "Loci from orthogroups to alignments. Each protein was aligned with MAFFT (--auto) and "
            "trimmed with trimAl (-automated1), and each coding sequence was read in the frame "
            "whose translation matches its protein, stop codons removed, then placed codon by "
            "codon on the aligned residues, keeping the columns trimAl kept. (b, c) Maximum "
            f"likelihood tree of the concatenated trimmed protein alignments (IQ-TREE "
            f"{version('iqtree', '3.1.3')}, one partition per locus, ModelFinder, 1,000 ultrafast "
            f"bootstrap replicates), rooted with {OG}, with gene (gCF) and site (sCF) "
            "concordance factors of each internal branch against the gene trees of the same loci; "
            "each branch is named by the ingroup clade it defines. (d) Rooted ingroup topology of "
            "each gene tree after rooting with the outgroup; resolved gene trees have both "
            "internal branches longer than IQ-TREE's minimum length, and the last columns count "
            "gene trees with both internal branches at UFBoot 95 or more. (e) D(P1, P2; P3, O) = "
            "(ABBA - BABA) / (ABBA + BABA) over codon alignment columns where all five taxa carry "
            "A, C, G or T, with the outgroup allele taken as ancestral: ABBA, P2 and P3 share the "
            "derived allele; BABA, P1 and P3 do; BBAA, P1 and P2 do. D is 0 when incomplete "
            "lineage sorting alone explains the discordant sites. Standard errors are from a "
            f"weighted block jackknife over {window} windows of the Cx. quinquefasciatus "
            "assembly, each locus placed by its reference gene; P is two sided from Z = D / SE, "
            "and Bonferroni P multiplies it by the number of tests. The tests were set before the "
            "run from the four taxon species tree; intact models as in Supp. Table S6(f).",
            "results/phylo/rooted/sco5_accounting.tsv; results/phylo/rooted/alignment_summary.tsv; "
            "results/phylo/rooted/rooted_summary.tsv; results/phylo/rooted/rooted_topology_counts.tsv; "
            "results/phylo/dstat/d_statistics.tsv",
            blocks)


# ------------------------------------------------------------------ build
def main():
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    sheets = [s1(), s2(), s3(), s4(), s5(), s6(), s7(), s8(), s9(), s10(), s11(), s12()]
    sheets = [x for x in sheets if x is not None]
    contents = []
    for name, title, caption, source, blocks in sheets:
        ws, nrows = write_sheet(wb, name, title, caption, source, blocks,
                                freeze=len(blocks) == 1)
        m = re.match(r"^Supp\. Table (S\d+)\. (.*)$", title)
        contents.append((m.group(1), name, m.group(2), nrows))
    fix_value_formats(wb, sheets)

    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 22
    readme.column_dimensions["C"].width = 80
    readme.column_dimensions["D"].width = 14
    readme["A1"] = "Supplementary Tables"
    readme["A1"].font = F_TITLE
    readme["A2"] = ("Phylogenomic relationships and gene content of four forms of the Culex "
                    "pipiens L. (Diptera: Culicidae) complex")
    readme["A3"] = "Tyler Maire, Erin Maley, Theresa Miller, Kyle Kosinski"
    for c in ("A2", "A3"):
        readme[c].font = F_BODY
    r = 5
    for i, h in enumerate(["Table", "Sheet", "Title", "Rows of data"], 1):
        c = readme.cell(row=r, column=i, value=h)
        c.font = F_BOLD
        c.border = Border(top=THIN, bottom=THIN)
    for t in contents:
        r += 1
        for i, v in enumerate(t, 1):
            c = readme.cell(row=r, column=i, value=v)
            c.font = F_BODY
            if i == 4:
                c.number_format = "#,##0"
    for i in range(1, 5):
        readme.cell(row=r, column=i).border = Border(bottom=THIN)
    og_acc, og_asm = ACCESSION[OUTGROUP]
    if OUTGROUP == "Cx_tarsalis":
        og_ref, models = f"{og_asm}, {og_acc}", ("Gene models of every form except quinquefasciatus "
                                                 "were transferred")
    else:
        og_ref = f"{og_acc}, {og_asm}, with its own Ensembl gene set"
        models = ("Gene models of pallens, molestus and pipiens were transferred")
    notes = [
        "Forms: quinquefasciatus, Cx. quinquefasciatus (GCF_015732765.1, source of the reference "
        "annotation); pallens, Cx. pipiens pallens (GCF_016801865.2); molestus, Cx. pipiens form "
        "molestus (GCA_024516115.1); pipiens, Cx. pipiens form pipiens (GCA_963924435.1); "
        f"{OG}, outgroup ({og_ref}).",
        f"{models} from the Cx. quinquefasciatus RefSeq annotation (NCBI Annotation Release 100) "
        f"with Liftoff {version('liftoff', '1.6.3')}, so their gene and transcript identifiers are "
        "those of the reference.",
        "Every value was produced by the workflow at https://github.com/tylermaire/cx_pipiens_pangenome "
        f"(release {RELEASE}); each sheet names the output files it was taken from.",
        "Transcript identifiers are RefSeq transcript accessions with the prefix rna-; the workflow's "
        "protein files drop the period before the version number (rna-XM_0382506301 for "
        "rna-XM_038250630.1).",
        "Percentages are on a 0 to 100 scale and labeled (%); shares, identities and retention are "
        "on a 0 to 1 scale. P values are two sided unless stated.",
    ]
    r += 2
    readme.cell(row=r, column=1, value="Notes").font = F_BOLD
    for n in notes:
        r += 1
        readme.cell(row=r, column=1, value=n).font = F_BODY
        readme.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        readme.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        readme.row_dimensions[r].height = 13.0 * math.ceil(len(n) / 125) + 3
    wb.save(OUT)
    print("wrote", OUT)
    for t in contents:
        print("  ", t)


def fix_value_formats(wb, sheets):
    """Blocks with a Measure | Value layout carry a format per row."""
    for name, title, caption, source, blocks in sheets:
        ws = wb[name]
        for b in blocks:
            fmts = getattr(b, "row_formats", None)
            if not fmts:
                continue
            # find the header row of this block by its first header and title
            start = None
            for row in ws.iter_rows(min_col=1, max_col=1):
                c = row[0]
                if c.value == b.title:
                    start = c.row + 2
                    break
            assert start, b.title
            for i, f in enumerate(fmts):
                c = ws.cell(row=start + i, column=2)
                raw = b.rows[i]["v"]
                v = cell_value(raw, "text" if f == "text" else f)
                c.value = v
                c.alignment = Alignment(horizontal="right")
                if isinstance(v, (int, float)) and f in NUMFMT:
                    c.number_format = NUMFMT[f]


if __name__ == "__main__":
    main()
